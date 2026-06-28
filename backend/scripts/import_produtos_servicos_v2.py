#!/usr/bin/env python3
"""Importa produtos e serviços faltantes da planilha Produtos_Servicos_v2.xlsx.

Compara o conteúdo da planilha com o que já existe no schema do tenant e cria
APENAS o que falta (idempotente):
  - produtos que ainda não existem (casa por nome normalizado, com aliases);
  - serviços que faltam em produtos (existentes ou recém-criados).

A coluna "PO (Responsável)" é mapeada para team_persons.full_name → responsavel_person_id.

Uso (dentro do container api):
  python scripts/import_produtos_servicos_v2.py /tmp/Produtos_Servicos_v2.xlsx --schema tenant_ss --dry-run
  python scripts/import_produtos_servicos_v2.py /tmp/Produtos_Servicos_v2.xlsx --schema tenant_ss
"""

from __future__ import annotations

import argparse
import asyncio
import re
import sys
import unicodedata
from collections import OrderedDict
from pathlib import Path

import openpyxl
from sqlalchemy import select, text

# app no PYTHONPATH (/app no container)
from app.core.database import AsyncSessionLocal
from app.modules.produtos.models import Product, ProductServico
from app.modules.teamops.models import Person


def _norm(s: str) -> str:
    s = unicodedata.normalize("NFD", s or "")
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", s.strip().lower())


# nome na planilha (normalizado) -> nome existente no banco (normalizado).
# "Solução 360" no banco está cadastrado com a grafia "Solicao 360".
PRODUCT_ALIASES: dict[str, str] = {
    "solucao 360": "solicao 360",
}

# placeholders da planilha que NÃO devem virar produto.
SKIP_PRODUCTS: set[str] = {"nao informado"}


def parse_workbook(path: Path) -> "OrderedDict[str, dict]":
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active  # planilha "Unificada"
    prods: "OrderedDict[str, dict]" = OrderedDict()
    for r in ws.iter_rows(min_row=2, values_only=True):
        prod = (str(r[0]).strip() if r[0] else "")
        if not prod:
            continue
        desc = (str(r[1]).strip() if len(r) > 1 and r[1] else "")
        po = (str(r[2]).strip() if len(r) > 2 and r[2] else "")
        serv = (str(r[3]).strip() if len(r) > 3 and r[3] else "")
        sdesc = (str(r[4]).strip() if len(r) > 4 and r[4] else "")
        e = prods.setdefault(prod, {"desc": "", "po": "", "servicos": OrderedDict()})
        if desc and not e["desc"]:
            e["desc"] = desc
        if po and not e["po"]:
            e["po"] = po
        if serv:
            # dedupe serviços repetidos dentro do mesmo produto (1ª descrição vence)
            e["servicos"].setdefault(_norm(serv), (serv, sdesc))
    wb.close()
    return prods


async def run(xlsx: Path, schema: str, ano: int, dry_run: bool) -> None:
    prods = parse_workbook(xlsx)
    print(f"Planilha: {xlsx}")
    print(f"Tenant schema: {schema}  |  ano_referencia: {ano}  |  dry_run: {dry_run}")
    print(f"Produtos na planilha: {len(prods)}\n")

    async with AsyncSessionLocal() as db:
        await db.execute(text(f"SET search_path TO {schema}, public"))

        persons = (await db.execute(select(Person))).scalars().all()
        person_by_name = {_norm(p.full_name): p for p in persons}

        products = (await db.execute(select(Product))).scalars().all()
        prod_by_key: dict[str, Product] = {_norm(p.name): p for p in products}

        created_products: list[str] = []
        created_services: list[tuple[str, str]] = []
        skipped: list[str] = []
        unmatched_po: set[str] = set()

        for excel_name, info in prods.items():
            key = _norm(excel_name)
            if key in SKIP_PRODUCTS:
                skipped.append(excel_name)
                continue
            target_key = PRODUCT_ALIASES.get(key, key)
            product = prod_by_key.get(target_key)

            if product is None:
                person = person_by_name.get(_norm(info["po"])) if info["po"] else None
                if info["po"] and person is None:
                    unmatched_po.add(info["po"])
                print(f"+ Produto NOVO: {excel_name}"
                      + (f"  [PO: {person.full_name}]" if person else ""))
                product = Product(
                    name=excel_name,
                    description=info["desc"] or None,
                    responsavel_person_id=person.id if person else None,
                )
                if not dry_run:
                    db.add(product)
                    await db.flush()
                prod_by_key[target_key] = product
                created_products.append(excel_name)
                existing_serv_keys: set[str] = set()
                next_order = 0
            else:
                serv_rows = (await db.execute(
                    select(ProductServico).where(ProductServico.product_id == product.id)
                )).scalars().all()
                existing_serv_keys = {_norm(s.name) for s in serv_rows}
                next_order = (max((s.order for s in serv_rows), default=-1)) + 1

            for skey, (sname, sdesc) in info["servicos"].items():
                if skey in existing_serv_keys:
                    continue
                print(f"    → serviço: {excel_name} :: {sname}")
                if not dry_run:
                    db.add(ProductServico(
                        product_id=product.id,
                        name=sname,
                        description=sdesc or None,
                        ano_referencia=ano,
                        order=next_order,
                        is_active=True,
                    ))
                existing_serv_keys.add(skey)
                created_services.append((excel_name, sname))
                next_order += 1

        if not dry_run:
            await db.commit()

        print("\n" + "=" * 60)
        print(f"Produtos criados:  {len(created_products)}")
        print(f"Serviços criados:  {len(created_services)}")
        print(f"Produtos ignorados (placeholder): {skipped}")
        if unmatched_po:
            print(f"POs sem match em team_persons: {sorted(unmatched_po)}")
        if dry_run:
            print("\n(DRY-RUN — nada foi gravado)")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx")
    parser.add_argument("--schema", default="tenant_ss")
    parser.add_argument("--ano", type=int, default=2026)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()
    path = Path(args.xlsx)
    if not path.is_file():
        sys.exit(f"Arquivo não encontrado: {path}")
    asyncio.run(run(path, args.schema, args.ano, args.dry_run))


if __name__ == "__main__":
    main()
